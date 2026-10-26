# -*- coding: utf-8 -*-
"""«Cargar nueva capa» of the energy module: the shared import window asks only for the file, its CRS and the layer
name, and creates a new layer in the project CRS or fails without creating any."""
import os
import tempfile
import unittest
import zipfile
from unittest import mock

from openpyxl import Workbook

from qgis.PyQt import QtCore, QtWidgets

_SETTINGS = tempfile.TemporaryDirectory()
QtCore.QSettings.setPath(QtCore.QSettings.NativeFormat, QtCore.QSettings.UserScope, _SETTINGS.name)
QtCore.QSettings.setPath(QtCore.QSettings.IniFormat, QtCore.QSettings.UserScope, _SETTINGS.name)
_APP = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])

from qgis.core import QgsCoordinateReferenceSystem, QgsCoordinateTransform, QgsPointXY, QgsProject  # noqa: E402

from VelantisWind import i18n  # noqa: E402
from VelantisWind.ui_core import layout_sources  # noqa: E402
from VelantisWind.ui_core.layout_sources import TurbineLayoutImportDialog, import_energy_layout  # noqa: E402

_CWD = os.getcwd()
_LANGS = ("es", "en", "fr", "de")


def setUpModule():
    # Without an initialised QGIS profile, QGIS writes its style database (symbology-style.db) in the working folder.
    os.chdir(_SETTINGS.name)


def tearDownModule():
    os.chdir(_CWD)


def _form_labels(dialog):
    form = dialog.layout().itemAt(0).layout()
    labels = []
    for row in range(form.rowCount()):
        item = form.itemAt(row, QtWidgets.QFormLayout.LabelRole)
        if item is not None and item.widget() is not None:
            labels.append(item.widget().text())
    return labels


def _in(lang, fn, *args, **kwargs):
    with mock.patch.object(i18n, "current_language", return_value=lang):
        return fn(*args, **kwargs)


LONLAT = [(-3.0, 40.0), (-2.99, 40.01)]
# A projected CRS in metres without an authority id.
CUSTOM = "+proj=tmerc +lat_0=0 +lon_0=-3.3 +k=0.9996 +x_0=500000 +y_0=0 +ellps=WGS84 +units=m +no_defs"


class ImportWindowTest(unittest.TestCase):
    def setUp(self):
        project = QgsProject.instance()
        project.clear()
        project.setCrs(QgsCoordinateReferenceSystem("EPSG:32630"))
        self.addCleanup(project.clear)

    def window(self, module):
        dialog = TurbineLayoutImportDialog(module=module)
        self.addCleanup(dialog.deleteLater)
        return dialog

    def test_energy_asks_only_for_the_name_the_file_and_its_crs(self):
        dialog = self.window("energy")
        self.assertEqual(_form_labels(dialog), ["Nombre de la capa:", "Fichero de coordenadas:", "Sistema de coordenadas:"])
        texts = [label.text() for label in dialog.findChildren(QtWidgets.QLabel)]
        self.assertFalse([t for t in texts if "en su tabla" in t])

    def test_noise_and_shadow_keep_their_fields(self):
        common = ["Nombre de la capa:", "Modelo / familia de aerogenerador:", "Parque eólico:", "Fichero de coordenadas:",
                  "Sistema de coordenadas:", "Altura buje:", "Diámetro rotor:", "Separación longitudinal:",
                  "Separación transversal:"]
        self.assertEqual(_form_labels(self.window("shadow")), common)
        self.assertEqual(_form_labels(self.window("noise")), common + ["Grupo fuente de ruido:", "LwA fijo por defecto:"])
        self.assertTrue([t for t in (l.text() for l in self.window("noise").findChildren(QtWidgets.QLabel)) if "en su tabla" in t])

    def test_energy_names_the_layer_after_the_file(self):
        dialog = self.window("energy")
        with mock.patch.object(QtWidgets.QFileDialog, "getOpenFileName", return_value=("/data/layout_a.csv", "")):
            dialog._browse_csv()
        self.assertEqual(dialog.ed_layer_name.text(), "layout_a")
        # Also when the path is typed instead of browsed.
        typed = self.window("energy")
        typed.ed_csv.setText("/data/layout_b.txt")
        self.assertEqual(typed.values()["layer_name"], "layout_b")

    def test_invalid_selected_crs_is_rejected(self):
        dialog = self.window("energy")
        if dialog.crs_widget is None:
            self.skipTest("QgsProjectionSelectionWidget is unavailable")
        dialog.crs_widget.setCrs(QgsCoordinateReferenceSystem())
        with self.assertRaises(ValueError):
            dialog.values()

    def test_texts_seen_from_energy_follow_the_language(self):
        spanish = {
            "title": "Importar layout de turbinas",
            "labels": ["Nombre de la capa:", "Fichero de coordenadas:", "Sistema de coordenadas:"],
            "placeholders": ["Ejemplo: Proyecto A · layout de aerogeneradores", "CSV/XLSX con columnas X,Y o easting,northing"],
            "button": "Examinar…",
        }
        for lang in _LANGS[1:]:
            with self.subTest(lang=lang):
                dialog = _in(lang, self.window, "energy")
                table = i18n._STATE["full"][lang]
                self.assertEqual(dialog.windowTitle(), table[spanish["title"]])
                self.assertEqual(_form_labels(dialog), [table[t] for t in spanish["labels"]])
                self.assertEqual([dialog.ed_layer_name.placeholderText(), dialog.ed_csv.placeholderText()],
                                 [table[t] for t in spanish["placeholders"]])
                self.assertIn(table[spanish["button"]], [b.text() for b in dialog.findChildren(QtWidgets.QPushButton)])
                seen = {}
                with mock.patch.object(QtWidgets.QFileDialog, "getOpenFileName",
                                       side_effect=lambda *a: seen.update(title=a[1], filters=a[3]) or ("", "")):
                    _in(lang, dialog._browse_csv)
                self.assertEqual(seen["title"], table["Selecciona el fichero de coordenadas de las turbinas"])
                self.assertEqual(seen["filters"], table["Ficheros de coordenadas (*.csv *.txt *.xlsx);;CSV/TXT (*.csv *.txt);;Libro de Excel (*.xlsx);;Todos los ficheros (*.*)"])


class ImportEnergyLayoutTest(unittest.TestCase):
    def setUp(self):
        QgsProject.instance().clear()
        self.addCleanup(QgsProject.instance().clear)
        QgsProject.instance().setCrs(QgsCoordinateReferenceSystem("EPSG:32630"))
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.folder = folder.name

    def write(self, name, text):
        path = os.path.join(self.folder, name)
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(text)
        return path

    def import_(self, path, crs="EPSG:32630", name=""):
        def accept(dialog):
            dialog.ed_csv.setText(path)
            dialog.ed_layer_name.setText(name)
            chosen = crs if isinstance(crs, QgsCoordinateReferenceSystem) else QgsCoordinateReferenceSystem(crs)
            if dialog.crs_widget is not None:
                dialog.crs_widget.setCrs(chosen)
            elif chosen.isValid():
                dialog.ed_crs.setText(chosen.authid() or chosen.toWkt(QgsCoordinateReferenceSystem.WKT_PREFERRED))
            else:
                dialog.ed_crs.setText("")
            return QtWidgets.QDialog.Accepted

        with mock.patch.object(layout_sources, "exec_dialog", side_effect=accept):
            return import_energy_layout()

    def turbine_layers(self):
        return [l for l in QgsProject.instance().mapLayers().values() if l.customProperty("velantis/layer_role", "") == "turbine_layout"]

    def points(self, layer):
        return [(f.geometry().asPoint().x(), f.geometry().asPoint().y()) for f in layer.getFeatures()]

    def expected(self, points, crs):
        transform = QgsCoordinateTransform(crs, QgsProject.instance().crs(), QgsProject.instance())
        return [(p.x(), p.y()) for p in (transform.transform(QgsPointXY(x, y)) for x, y in points)]

    def assertPointsAlmostEqual(self, got, expected):
        self.assertEqual(len(got), len(expected))
        for (x, y), (ex, ey) in zip(got, expected):
            self.assertAlmostEqual(x, ex, delta=1e-6)
            self.assertAlmostEqual(y, ey, delta=1e-6)

    def test_points_are_transformed_from_the_chosen_crs_to_the_project(self):
        path = self.write("layout_a.csv", "lon,lat\n" + "".join(f"{x},{y}\n" for x, y in LONLAT))
        layer = self.import_(path, crs="EPSG:4326")
        self.assertEqual(layer.name(), "layout_a")
        self.assertEqual(layer.crs(), QgsProject.instance().crs())
        self.assertPointsAlmostEqual(self.points(layer), self.expected(LONLAT, QgsCoordinateReferenceSystem("EPSG:4326")))
        # No turbine data: the row writes them when the layer is selected.
        self.assertEqual(layer.customProperty("velantis/model_name", ""), "")
        self.assertIsNone(layer.customProperty("velantis/hub_height_m", None))
        self.assertIsNone(layer.customProperty("velantis/diameter_m", None))

    def test_crs_without_authority_id_is_kept(self):
        # Also a derived projected CRS (UTM 30N shifted 100 m), which has no WKT1: it is not taken as the project's.
        from tests.test_energy_project_crs import DERIVED_PROJECTED
        path = self.write("layout.csv", "x,y\n500000,4400000\n500300,4400100\n")
        for custom in (QgsCoordinateReferenceSystem.fromProj(CUSTOM), QgsCoordinateReferenceSystem.fromWkt(DERIVED_PROJECTED)):
            with self.subTest(crs=custom.description()):
                self.assertTrue(custom.isValid())
                self.assertEqual(custom.authid(), "")
                layer = self.import_(path, crs=custom)
                self.assertPointsAlmostEqual(self.points(layer), self.expected([(500000.0, 4400000.0), (500300.0, 4400100.0)], custom))
                self.assertNotEqual(self.points(layer)[0], (500000.0, 4400000.0))

    def test_without_the_crs_selector_the_project_crs_is_proposed_whole(self):
        # Where QGIS has no CRS selector the window proposes the project's CRS as text: also a derived one, without WKT1.
        from tests.test_energy_project_crs import DERIVED_PROJECTED
        QgsProject.instance().setCrs(QgsCoordinateReferenceSystem.fromWkt(DERIVED_PROJECTED))
        with mock.patch.object(layout_sources, "QgsProjectionSelectionWidget", None):
            dialog = TurbineLayoutImportDialog(module="energy")
        self.addCleanup(dialog.deleteLater)
        self.assertIsNone(dialog.crs_widget)
        self.assertEqual(QgsCoordinateReferenceSystem(dialog.ed_crs.text()), QgsProject.instance().crs())

    def test_project_with_derived_crs_keeps_layer_crs_and_metadata(self):
        from tests.test_energy_project_crs import DERIVED_PROJECTED
        target = QgsCoordinateReferenceSystem.fromWkt(DERIVED_PROJECTED)
        self.assertTrue(target.isValid())
        self.assertFalse(target.authid())
        self.assertFalse(target.toWkt())
        QgsProject.instance().setCrs(target)
        path = self.write("derived_target.csv", "x,y\n500000,4400000\n")
        layer = self.import_(path, crs="EPSG:32630")
        self.assertTrue(layer.isValid())
        self.assertTrue(layer.crs().isValid())
        self.assertEqual(layer.crs(), target)
        self.assertEqual(
            layer.customProperty("velantis/layout_crs"),
            target.toWkt(QgsCoordinateReferenceSystem.WKT_PREFERRED),
        )
        self.assertEqual(layer.customProperty("velantis/source_crs"), "EPSG:32630")
        self.assertPointsAlmostEqual(
            self.points(layer),
            self.expected([(500000.0, 4400000.0)], QgsCoordinateReferenceSystem("EPSG:32630")),
        )

    def test_noise_and_shadow_import_custom_or_derived_source_crs(self):
        from tests.test_energy_project_crs import DERIVED_PROJECTED
        target = QgsProject.instance().crs()
        sources = (QgsCoordinateReferenceSystem.fromProj(CUSTOM),
                   QgsCoordinateReferenceSystem.fromWkt(DERIVED_PROJECTED))
        path = self.write("other_modules.csv", "x,y\n500000,4400000\n")
        for module in ("noise", "shadow"):
            for chosen in sources:
                with self.subTest(module=module, crs=chosen.description()):
                    def accept(dialog):
                        dialog.ed_csv.setText(path)
                        if dialog.crs_widget is not None:
                            dialog.crs_widget.setCrs(chosen)
                        else:
                            dialog.ed_crs.setText(chosen.toWkt(QgsCoordinateReferenceSystem.WKT_PREFERRED))
                        return QtWidgets.QDialog.Accepted

                    with mock.patch.object(layout_sources, "exec_dialog", side_effect=accept):
                        layer = layout_sources.import_turbine_layout_from_csv(module=module)
                    self.assertTrue(layer.isValid())
                    self.assertEqual(layer.crs(), target)
                    self.assertEqual(
                        layer.customProperty("velantis/source_crs"),
                        chosen.toWkt(QgsCoordinateReferenceSystem.WKT_PREFERRED),
                    )
                    self.assertPointsAlmostEqual(
                        self.points(layer), self.expected([(500000.0, 4400000.0)], chosen)
                    )

    def test_noise_and_shadow_no_selector_keeps_derived_project_crs(self):
        from tests.test_energy_project_crs import DERIVED_PROJECTED
        target = QgsCoordinateReferenceSystem.fromWkt(DERIVED_PROJECTED)
        QgsProject.instance().setCrs(target)
        with mock.patch.object(layout_sources, "QgsProjectionSelectionWidget", None):
            for module in ("noise", "shadow"):
                dialog = TurbineLayoutImportDialog(module=module)
                self.addCleanup(dialog.deleteLater)
                self.assertEqual(QgsCoordinateReferenceSystem(dialog.ed_crs.text()), target)

    def test_invalid_crs_is_rejected_without_a_layer(self):
        path = self.write("layout.csv", "x,y\n500000,4400000\n")
        with self.assertRaises(ValueError):
            self.import_(path, crs=QgsCoordinateReferenceSystem())
        self.assertEqual(QgsProject.instance().mapLayers(), {})

    def test_importing_twice_creates_two_layers(self):
        path = self.write("layout_a.csv", "x,y\n500000,4400000\n")
        first, second = self.import_(path), self.import_(path)
        self.assertNotEqual(first.id(), second.id())
        self.assertNotEqual(first.name(), second.name())
        self.assertEqual(len(self.turbine_layers()), 2)

    def test_file_without_coordinates_fails_in_every_language_without_a_layer(self):
        path = self.write("empty.csv", "x,y\n")
        messages = {}
        for lang in _LANGS:
            with self.assertRaises(ValueError) as caught:
                _in(lang, self.import_, path)
            messages[lang] = str(caught.exception)
        self.assertEqual(QgsProject.instance().mapLayers(), {})
        reason = "No se han encontrado coordenadas XY válidas. Usa columnas llamadas X/Y, easting/northing o longitude/latitude, o un fichero de dos columnas."
        context = "No se pudo leer el fichero de coordenadas: {error}"
        self.assertEqual(messages["es"], context.format(error=reason))
        for lang in _LANGS[1:]:
            table = i18n._STATE["full"][lang]
            self.assertEqual(messages[lang], table[context].format(error=table[reason]))

    def test_unreadable_files_fail_in_every_language_without_a_layer(self):
        os.mkdir(os.path.join(self.folder, "folder.csv"))
        with open(os.path.join(self.folder, "latin1.csv"), "wb") as fh:
            fh.write("x,y\n500000,4400000\né\xff\n".encode("latin-1"))
        broken = self.write("broken.xlsx", "not a workbook")
        # A valid workbook whose numeric cell holds text: openpyxl raises its own ValueError while reading it.
        bad_cell = os.path.join(self.folder, "bad_cell.xlsx")
        book = Workbook()
        book.active.append(("x", "y"))
        book.active.append((500000, 4400000))
        book.save(bad_cell)
        with zipfile.ZipFile(bad_cell) as zin:
            items = {name: zin.read(name) for name in zin.namelist()}
        sheet = "xl/worksheets/sheet1.xml"
        items[sheet] = items[sheet].replace(b"<v>500000</v>", b"<v>invalid-number</v>")
        with zipfile.ZipFile(bad_cell, "w") as zout:
            for name, data in items.items():
                zout.writestr(name, data)
        for path in (os.path.join(self.folder, "folder.csv"), os.path.join(self.folder, "latin1.csv"), broken, bad_cell):
            for lang in _LANGS:
                with self.subTest(path=os.path.basename(path), lang=lang):
                    with self.assertRaises(ValueError) as caught:
                        _in(lang, self.import_, path)
                    template = "No se pudo leer el fichero de coordenadas: {error}"
                    prefix = (template if lang == "es" else i18n._STATE["full"][lang][template]).split("{error}")[0]
                    self.assertTrue(str(caught.exception).startswith(prefix), str(caught.exception))
        self.assertEqual(self.turbine_layers(), [])

    def test_without_the_crs_selector_the_typed_crs_is_kept(self):
        with mock.patch.object(layout_sources, "QgsProjectionSelectionWidget", None):
            dialog = TurbineLayoutImportDialog(module="energy")
            self.addCleanup(dialog.deleteLater)
            self.assertEqual(dialog.ed_crs.text(), "EPSG:32630")
            self.assertEqual(_form_labels(dialog), ["Nombre de la capa:", "Fichero de coordenadas:", "Sistema de coordenadas:"])
            french = _in("fr", TurbineLayoutImportDialog, module="energy")
            self.addCleanup(french.deleteLater)
            self.assertEqual(french.ed_crs.placeholderText(), i18n._STATE["full"]["fr"]["Ejemplo: EPSG:4326"])
            path = self.write("layout.csv", "x,y\n500000,4400000\n500300,4400100\n")
            custom = QgsCoordinateReferenceSystem.fromProj(CUSTOM)
            layer = self.import_(path, crs=custom)
            self.assertPointsAlmostEqual(self.points(layer), self.expected([(500000.0, 4400000.0), (500300.0, 4400100.0)], custom))
            with self.assertRaises(ValueError):
                self.import_(path, crs=QgsCoordinateReferenceSystem())
        self.assertEqual(len(self.turbine_layers()), 1)

    def test_cancel_creates_nothing(self):
        with mock.patch.object(layout_sources, "exec_dialog", return_value=QtWidgets.QDialog.Rejected):
            self.assertIsNone(import_energy_layout())
        self.assertEqual(QgsProject.instance().mapLayers(), {})


if __name__ == "__main__":
    unittest.main()
