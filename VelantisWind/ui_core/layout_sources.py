# -*- coding: utf-8 -*-
"""Shared turbine-layout import helpers for VelantisWind modules.

The Energy, Noise and Shadow modules all need the same first step: obtain a
point layer with turbine coordinates and enough metadata for the calculation
module to interpret it.  Keeping the CSV parsing and memory-layer creation here
prevents Noise/Shadow from depending on the Energy dialog just to load a layout.
"""
from __future__ import annotations

import logging

import csv
import math

from ..ag_core.curve_text import autodetected_table_rows
import os
from typing import Dict, List, Optional, Tuple

from qgis.PyQt import QtCore, QtWidgets

try:
    from qgis.gui import QgsProjectionSelectionWidget
except Exception:  # pragma: no cover - unusual stripped QGIS builds
    QgsProjectionSelectionWidget = None

try:
    from ..i18n import apply_i18n, tr_text
except Exception:  # pragma: no cover
    def apply_i18n(_widget):
        return None
    def tr_text(text):
        return text

from qgis.core import (
    QgsCoordinateReferenceSystem,
    QgsCoordinateTransform,
    QgsFeature,
    QgsField,
    QgsGeometry,
    QgsPointXY,
    QgsProject,
    QgsVectorLayer,
)

_LAYOUT_GROUP_NAME = "VelantisWind · Turbine layouts"
# Projected names first: a table with both pairs is read in its projected columns, and the axes of the two pairs are
# never mixed.
_XY_NAMES = (
    ({"x", "utm_x", "easting", "east"}, {"y", "utm_y", "northing", "north"}),
    ({"lon", "longitude"}, {"lat", "latitude"}),
)


def _safe_float(value) -> Optional[float]:
    try:
        number = float(str(value).strip().replace(",", "."))
    except Exception:
        return None
    # NaN, inf or 1e999 read as numbers, but are no coordinate.
    return number if math.isfinite(number) else None


def _reads_as_number(cells) -> bool:
    """True if any cell reads as a number, finite or not: such a row is data, not a header."""
    for cell in cells:
        try:
            float(str(cell).strip().replace(",", "."))
            return True
        except Exception:  # nosec B112
            continue
    return False


def _rows_to_points(rows, lines: Optional[List[int]] = None) -> List[Tuple[float, float]]:
    # An empty XLSX cell is None, whose text «None» is not empty: an empty first row would take the header's place.
    # Each row keeps its number in the file (``lines``: the line where each record starts), to name the ones that
    # cannot be read.
    numbered = [
        (line, list(r)) for line, r in (zip(lines, rows) if lines is not None else enumerate(rows, start=1))
        if r and any(c is not None and str(c).strip() for c in r)
    ]
    if not numbered:
        raise ValueError(tr_text("El fichero de coordenadas está vacío."))
    rows = [r for _line, r in numbered]

    header = [str(c).strip().lower() for c in rows[0]]
    x_idx = y_idx = None
    for x_names, y_names in _XY_NAMES:
        x_idx = next((i for i, h in enumerate(header) if h in x_names), None)
        y_idx = next((i for i, h in enumerate(header) if h in y_names), None)
        if x_idx is not None and y_idx is not None:
            break

    named = x_idx is not None and y_idx is not None
    data_rows = numbered[1:] if named else numbered
    if not named:
        # Names of both pairs and neither complete (lon and northing): the first two columns would mix two systems.
        if all(any(h in x_names | y_names for h in header) for x_names, y_names in _XY_NAMES):
            data_rows = []
        x_idx, y_idx = 0, 1

    points: List[Tuple[float, float]] = []
    unreadable: List[int] = []
    for n, (line, row) in enumerate(data_rows):
        x = _safe_float(row[x_idx]) if len(row) > x_idx else None
        y = _safe_float(row[y_idx]) if len(row) > y_idx else None
        if x is None or y is None:
            # Without named columns the first row may be a header of other names, if neither coordinate reads as a
            # number; any other row is a turbine whose coordinates cannot be read, and a layout without it would be
            # computed without a word.
            if not (n == 0 and not named and not _reads_as_number(row[:2])):
                unreadable.append(line)
            continue
        points.append((float(x), float(y)))
    if unreadable and points:
        shown = ", ".join(str(line) for line in unreadable[:10]) + ("…" if len(unreadable) > 10 else "")
        raise ValueError(tr_text(
            "Fila(s) del fichero sin coordenadas X e Y válidas: {rows}. Corrígelas antes de importar."
        ).format(rows=shown))

    if not points:
        raise ValueError(tr_text(
            "No se han encontrado coordenadas XY válidas. Usa columnas llamadas X/Y, "
            "easting/northing o longitude/latitude, o un fichero de dos columnas."
        ))
    return points


def read_xy_table(path: str) -> List[Tuple[float, float]]:
    """Read turbine coordinates from CSV/TXT or the first XLSX sheet.

    XLSX support is optional and uses ``openpyxl`` when it is available in the
    QGIS Python environment. CSV/TXT remains dependency-free.
    """
    if not path or not os.path.exists(path):
        raise ValueError(tr_text("No se encuentra el fichero de coordenadas."))

    suffix = os.path.splitext(path)[1].lower()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            raise RuntimeError(tr_text(
                "Para importar un XLSX hace falta openpyxl en el Python de QGIS. "
                "Guarda el fichero como CSV o instala esa dependencia opcional."
            )) from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            return _rows_to_points(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()

    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t ")
        except Exception:
            dialect = csv.excel
            # Aligned or mixed-whitespace TXT can defeat CSV's delimiter sniffer.
            delimiter = "," if "," in sample else ";" if ";" in sample else " "
        else:
            delimiter = dialect.delimiter
        reader = autodetected_table_rows(f, delimiter, dialect=dialect)
        rows: List[List[str]] = []
        lines: List[int] = []
        start = 1
        for row in reader:
            # A quoted field may span several lines: the record starts after the previous one ends.
            rows.append(row)
            lines.append(start)
            start = reader.line_num + 1
        return _rows_to_points(rows, lines)


def read_xy_csv(path: str) -> List[Tuple[float, float]]:
    """Backward-compatible alias for coordinate-table import."""
    return read_xy_table(path)


def _crs_definition(crs: QgsCoordinateReferenceSystem) -> str:
    """A portable CRS definition: QGIS-local USER ids are not usable by GDAL."""
    authid = crs.authid()
    return authid if authid and not authid.upper().startswith("USER:") else crs.toWkt(
        QgsCoordinateReferenceSystem.WKT_PREFERRED
    )


def _project_crs_uri() -> str:
    crs = QgsProject.instance().crs()
    return _crs_definition(crs) if crs and crs.isValid() else "EPSG:25830"


def _validated_crs(authid: str = "") -> QgsCoordinateReferenceSystem:
    candidate = str(authid or "").strip()
    if candidate:
        crs = QgsCoordinateReferenceSystem(candidate)
        if crs.isValid():
            return crs
        raise ValueError(tr_text("Sistema de coordenadas no válido: {crs}").format(crs=candidate))
    project_crs = QgsProject.instance().crs()
    if project_crs and project_crs.isValid():
        return project_crs
    return QgsCoordinateReferenceSystem("EPSG:25830")


def _ensure_group(name: str = _LAYOUT_GROUP_NAME):
    root = QgsProject.instance().layerTreeRoot()
    try:
        group = root.findGroup(name)
    except Exception:
        group = None
    if group is None:
        group = root.addGroup(name)
    return group


def _unique_layer_name(base_name: str) -> str:
    existing = {lyr.name() for lyr in QgsProject.instance().mapLayers().values()}
    base = (base_name or "Turbine layout").strip() or "Turbine layout"
    if base not in existing:
        return base
    i = 2
    while True:
        candidate = f"{base} {i}"
        if candidate not in existing:
            return candidate
        i += 1


def create_turbine_layout_layer(
    *,
    coords: List[Tuple[float, float]],
    layer_name: str,
    module: str,
    model_name: str,
    hub_height_m: Optional[float] = None,
    rotor_diameter_m: Optional[float] = None,
    coords_csv: str = "",
    park_name: str = "",
    source_group_name: str = "",
    extra_properties: Optional[Dict[str, object]] = None,
    crs_authid: str = "",
) -> QgsVectorLayer:
    """Create a memory point layer and tag it as a Velantis turbine layout."""
    if not coords:
        raise ValueError(tr_text("No se han indicado coordenadas de turbinas."))

    layer_name = _unique_layer_name(layer_name)
    source_crs = _validated_crs(crs_authid)
    project_crs = QgsProject.instance().crs()
    target_crs = project_crs if project_crs and project_crs.isValid() else source_crs
    # A derived projected CRS can have neither an authid nor a WKT1 representation.
    # For those CRSes the memory-provider URI cannot reliably express the definition:
    # create the layer first and assign the full QGIS CRS object explicitly.
    layer_crs = target_crs.authid()
    if layer_crs.upper().startswith("USER:"):
        layer_crs = ""  # The provider must not resolve a QGIS-local authority id.
    lyr = QgsVectorLayer(f"Point?crs={layer_crs}" if layer_crs else "Point", layer_name, "memory")
    if not lyr.isValid():
        raise RuntimeError(tr_text("No se pudo crear la capa en memoria del layout de turbinas."))
    lyr.setCrs(target_crs)

    provider = lyr.dataProvider()
    fields = [
        QgsField("model", QtCore.QVariant.String),
        QgsField("park", QtCore.QVariant.String),
        QgsField("module", QtCore.QVariant.String),
        QgsField("hh_m", QtCore.QVariant.Double),
        QgsField("diam_m", QtCore.QVariant.Double),
        QgsField("source_csv", QtCore.QVariant.String),
    ]
    if module == "noise":
        fields.append(QgsField("lwa_dba", QtCore.QVariant.Double))
    provider.addAttributes(fields)
    lyr.updateFields()

    default_lwa = None
    if extra_properties and extra_properties.get("velantis/default_lwa_dba") is not None:
        default_lwa = _safe_float(extra_properties.get("velantis/default_lwa_dba"))

    coordinate_transform = None
    if source_crs != target_crs:
        coordinate_transform = QgsCoordinateTransform(
            source_crs, target_crs, QgsProject.instance().transformContext()
        )

    feats: List[QgsFeature] = []
    for x, y in coords:
        point = QgsPointXY(float(x), float(y))
        if coordinate_transform is not None:
            try:
                point = coordinate_transform.transform(point)
            except Exception as exc:
                raise RuntimeError(tr_text(
                    "No se pudieron transformar las coordenadas de las turbinas ({source} → {target}): {error}"
                ).format(
                    source=source_crs.authid() or tr_text("sistema de origen"),
                    target=target_crs.authid() or tr_text("sistema del proyecto"),
                    error=exc,
                )) from exc
        feat = QgsFeature(lyr.fields())
        feat.setGeometry(QgsGeometry.fromPointXY(point))
        attrs = [
            model_name,
            park_name,
            module,
            float(hub_height_m) if hub_height_m is not None else None,
            float(rotor_diameter_m) if rotor_diameter_m is not None else None,
            coords_csv,
        ]
        if module == "noise":
            attrs.append(float(default_lwa) if default_lwa is not None else None)
        feat.setAttributes(attrs)
        feats.append(feat)
    provider.addFeatures(feats)
    lyr.updateExtents()

    # Generic metadata used by all module detectors.
    lyr.setCustomProperty("velantis/layer_role", "turbine_layout")
    lyr.setCustomProperty("velantis/source_module", module)
    lyr.setCustomProperty("velantis/model_name", model_name)
    lyr.setCustomProperty("velantis/coords_csv", coords_csv)
    lyr.setCustomProperty("velantis/source_crs", _crs_definition(source_crs))
    lyr.setCustomProperty("velantis/layout_crs", _crs_definition(target_crs))
    if park_name:
        lyr.setCustomProperty("velantis/park_name", park_name)
    if source_group_name:
        lyr.setCustomProperty("velantis/noise_group_name", source_group_name)
    if hub_height_m is not None:
        lyr.setCustomProperty("velantis/hub_height_m", float(hub_height_m))
    if rotor_diameter_m is not None:
        lyr.setCustomProperty("velantis/diameter_m", float(rotor_diameter_m))
    for key, value in (extra_properties or {}).items():
        try:
            lyr.setCustomProperty(str(key), value)
        except Exception:  # nosec B110
            pass

    QgsProject.instance().addMapLayer(lyr, False)
    _ensure_group().addLayer(lyr)
    try:
        # Las de Energía las crea el diálogo después de elegir la capa en su fila, junto con las demás filas (#96).
        if module != "energy" and not lyr.crs().isGeographic():
            from ..spacing_core.auto_envelope import ensure_spacing_envelope_for_layer
            ensure_spacing_envelope_for_layer(lyr)
    except Exception:  # nosec B110
        pass
    return lyr


class TurbineLayoutImportDialog(QtWidgets.QDialog):
    """Small module-aware coordinate-table import dialog."""

    def __init__(self, parent=None, module: str = "noise"):
        super().__init__(parent)
        self.module = (module or "noise").strip().lower()
        # Energía solo pide el fichero, su sistema y el nombre de la capa: la turbina se define en su fila.
        energy = self.module == "energy"
        self.setWindowTitle("Importar layout de turbinas")
        self.setMinimumWidth(520)

        lay = QtWidgets.QVBoxLayout(self)
        form = QtWidgets.QFormLayout()

        self.ed_layer_name = QtWidgets.QLineEdit()
        self.ed_layer_name.setPlaceholderText("Ejemplo: Proyecto A · layout de aerogeneradores")
        form.addRow("Nombre de la capa:", self.ed_layer_name)

        self._default_model = tr_text("Modelo de aerogenerador")
        self.ed_model = QtWidgets.QLineEdit(self._default_model)
        if not energy:
            form.addRow("Modelo / familia de aerogenerador:", self.ed_model)

        self.ed_park = QtWidgets.QLineEdit()
        self.ed_park.setPlaceholderText("Opcional")
        if not energy:
            form.addRow("Parque eólico:", self.ed_park)

        self.ed_csv = QtWidgets.QLineEdit()
        self.ed_csv.setPlaceholderText("CSV/XLSX con columnas X,Y o easting,northing")
        btn_csv = QtWidgets.QPushButton("Examinar…")
        btn_csv.clicked.connect(self._browse_csv)
        csv_row = QtWidgets.QHBoxLayout()
        csv_row.addWidget(self.ed_csv, 1)
        csv_row.addWidget(btn_csv)
        form.addRow("Fichero de coordenadas:", csv_row)

        self.crs_widget = QgsProjectionSelectionWidget(self) if QgsProjectionSelectionWidget is not None else None
        if self.crs_widget is not None:
            try:
                self.crs_widget.setCrs(QgsProject.instance().crs())
            except Exception as _ignored_exc:
                logging.getLogger(__name__).debug("Ignored optional operation at %s:%d: %s", __name__, 393, _ignored_exc)
            form.addRow("Sistema de coordenadas:", self.crs_widget)
        else:
            project_crs = QgsProject.instance().crs()
            if energy:
                # Energía respeta el sistema del proyecto aunque no tenga identificador de autoridad (en WKT2: un
                # proyectado derivado no tiene WKT1).
                default_crs = (
                    _crs_definition(project_crs)
                ) if project_crs.isValid() else ""
            else:
                default_crs = _project_crs_uri()
            self.ed_crs = QtWidgets.QLineEdit(default_crs)
            self.ed_crs.setPlaceholderText("Ejemplo: EPSG:4326")
            form.addRow("Sistema de coordenadas:", self.ed_crs)

        self.sp_hh = QtWidgets.QDoubleSpinBox()
        self.sp_hh.setDecimals(1)
        self.sp_hh.setRange(0.0, 400.0)
        self.sp_hh.setValue(100.0)
        self.sp_hh.setSuffix(" m")
        if not energy:
            form.addRow("Altura buje:", self.sp_hh)

        self.sp_diam = QtWidgets.QDoubleSpinBox()
        self.sp_diam.setDecimals(1)
        self.sp_diam.setRange(0.0, 400.0)
        self.sp_diam.setValue(120.0)
        self.sp_diam.setSuffix(" m")
        if not energy:
            form.addRow("Diámetro rotor:", self.sp_diam)

        self.sp_spacing_long = QtWidgets.QDoubleSpinBox()
        self.sp_spacing_long.setRange(0.5, 30.0)
        self.sp_spacing_long.setDecimals(2)
        self.sp_spacing_long.setSingleStep(0.5)
        self.sp_spacing_long.setValue(7.0)
        self.sp_spacing_long.setSuffix(" · D")
        if not energy:
            form.addRow("Separación longitudinal:", self.sp_spacing_long)

        self.sp_spacing_trans = QtWidgets.QDoubleSpinBox()
        self.sp_spacing_trans.setRange(0.5, 30.0)
        self.sp_spacing_trans.setDecimals(2)
        self.sp_spacing_trans.setSingleStep(0.5)
        self.sp_spacing_trans.setValue(4.0)
        self.sp_spacing_trans.setSuffix(" · D")
        if not energy:
            form.addRow("Separación transversal:", self.sp_spacing_trans)

        self.ed_source_group = QtWidgets.QLineEdit(tr_text("Grupo fuente acústico") if self.module == "noise" else "")
        if self.module == "noise":
            form.addRow("Grupo fuente de ruido:", self.ed_source_group)

            self.sp_lwa = QtWidgets.QDoubleSpinBox()
            self.sp_lwa.setDecimals(1)
            self.sp_lwa.setRange(0.0, 200.0)
            self.sp_lwa.setValue(105.0)
            self.sp_lwa.setSuffix(" dB(A)")
            form.addRow("LwA fijo por defecto:", self.sp_lwa)
        else:
            self.sp_lwa = None

        lay.addLayout(form)

        # Energía no tiene la tabla del módulo de la que habla la nota.
        if not energy:
            note = QtWidgets.QLabel(
                "La capa se añadirá a QGIS como layout de turbinas de VelantisWind. "
                "Antes de calcular, aún puedes editar los valores propios del módulo en su tabla."
            )
            note.setWordWrap(True)
            lay.addWidget(note)

        buttons = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel)
        # The standard buttons follow the language of Qt, not the plugin's.
        buttons.button(QtWidgets.QDialogButtonBox.Ok).setText("Aceptar")
        buttons.button(QtWidgets.QDialogButtonBox.Cancel).setText("Cancelar")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        lay.addWidget(buttons)
        apply_i18n(self)

    def _browse_csv(self):
        start_dir = os.path.dirname(self.ed_csv.text().strip()) if self.ed_csv.text().strip() else os.path.expanduser("~")
        path, _ = QtWidgets.QFileDialog.getOpenFileName(
            self,
            tr_text("Selecciona el fichero de coordenadas de las turbinas"),
            start_dir,
            tr_text("Ficheros de coordenadas (*.csv *.txt *.xlsx);;CSV/TXT (*.csv *.txt);;Libro de Excel (*.xlsx);;Todos los ficheros (*.*)"),
        )
        if path:
            self.ed_csv.setText(path)
            if not self.ed_layer_name.text().strip():
                base = os.path.splitext(os.path.basename(path))[0]
                if self.module == "energy":
                    self.ed_layer_name.setText(base)
                else:
                    suffix = {"noise": "noise layout", "shadow": "shadow layout"}.get(self.module, "turbine layout")
                    self.ed_layer_name.setText(f"{base} · {suffix}")

    def values(self) -> Dict[str, object]:
        model = self.ed_model.text().strip() or self._default_model
        # Layer names are identifiers: keep the original default component independent of the UI language.
        layer_model = "WT model" if model == self._default_model else model
        layer_name = self.ed_layer_name.text().strip()
        if not layer_name and self.module == "energy":
            layer_name = os.path.splitext(os.path.basename(self.ed_csv.text().strip()))[0]
        layer_name = layer_name or f"{layer_model} · {self.module} layout"
        if self.crs_widget is not None:
            selected_crs = self.crs_widget.crs()
            if not selected_crs.isValid():
                raise ValueError(tr_text("Sistema de coordenadas no válido: {crs}").format(crs=""))
            crs_authid = _crs_definition(selected_crs)
        else:
            crs_authid = self.ed_crs.text().strip() or _project_crs_uri()

        out = {
            "layer_name": layer_name,
            "model_name": model,
            "park_name": self.ed_park.text().strip(),
            "coords_csv": self.ed_csv.text().strip(),
            "crs_authid": crs_authid,
            "hub_height_m": float(self.sp_hh.value()) if self.sp_hh.value() > 0 else None,
            "rotor_diameter_m": float(self.sp_diam.value()) if self.sp_diam.value() > 0 else None,
            "spacing_long_d": float(self.sp_spacing_long.value()),
            "spacing_trans_d": float(self.sp_spacing_trans.value()),
            "source_group_name": self.ed_source_group.text().strip() if self.module == "noise" else "",
        }
        if self.module == "noise" and self.sp_lwa is not None:
            out["default_lwa_dba"] = float(self.sp_lwa.value())
        return out


def exec_dialog(dialog: QtWidgets.QDialog) -> int:
    """Qt5/Qt6 compatible dialog execution."""
    if hasattr(dialog, "exec"):
        return dialog.exec()
    return dialog.exec_()


def import_turbine_layout_from_csv(parent=None, module: str = "noise") -> Optional[QgsVectorLayer]:
    """Prompt for CSV/TXT/XLSX coordinates and create a tagged QGIS layer."""
    dlg = TurbineLayoutImportDialog(parent, module=module)
    if exec_dialog(dlg) != QtWidgets.QDialog.Accepted:
        return None
    values = dlg.values()
    coords_csv = str(values.get("coords_csv") or "").strip()
    coords = read_xy_table(coords_csv)
    extra: Dict[str, object] = {}
    if module == "noise" and values.get("default_lwa_dba") is not None:
        extra["velantis/default_lwa_dba"] = float(values.get("default_lwa_dba"))
    extra["velantis/spacing_long_d"] = float(values.get("spacing_long_d") or 7.0)
    extra["velantis/spacing_trans_d"] = float(values.get("spacing_trans_d") or 4.0)
    extra["velantis/spacing_mode"] = "auto_energy"
    extra["velantis/spacing_angle_deg"] = 0.0
    return create_turbine_layout_layer(
        coords=coords,
        layer_name=str(values.get("layer_name") or "Turbine layout"),
        module=module,
        model_name=str(values.get("model_name") or tr_text("Modelo de aerogenerador")),
        hub_height_m=values.get("hub_height_m"),
        rotor_diameter_m=values.get("rotor_diameter_m"),
        coords_csv=coords_csv,
        park_name=str(values.get("park_name") or ""),
        source_group_name=str(values.get("source_group_name") or ""),
        extra_properties=extra,
        crs_authid=str(values.get("crs_authid") or ""),
    )


def import_energy_layout(parent=None) -> Optional[QgsVectorLayer]:
    """«Cargar nueva capa» de Energía: una capa nueva con las coordenadas de un fichero, sin datos de turbina.

    Devuelve ``None`` si el usuario cancela. Si el fichero, el sistema elegido o la transformación fallan, lanza
    ``ValueError`` o ``RuntimeError`` con el mensaje ya traducido y no crea ninguna capa.
    """
    dlg = TurbineLayoutImportDialog(parent, module="energy")
    if exec_dialog(dlg) != QtWidgets.QDialog.Accepted:
        return None
    values = dlg.values()
    if dlg.crs_widget is not None:
        crs = dlg.crs_widget.crs()
    else:
        text = dlg.ed_crs.text().strip()
        crs = QgsCoordinateReferenceSystem(text) if text else QgsCoordinateReferenceSystem()
    if not crs.isValid():
        raise ValueError(tr_text("El sistema de coordenadas elegido no es válido."))
    coords_csv = str(values.get("coords_csv") or "").strip()
    try:
        coords = read_xy_table(coords_csv)
    except Exception as exc:  # también un directorio, otra codificación o un XLSX dañado
        raise ValueError(tr_text("No se pudo leer el fichero de coordenadas: {error}").format(error=exc)) from exc
    # Un sistema sin identificador de autoridad se pasa completo, en WKT2 (un proyectado derivado no tiene WKT1): si no,
    # se tomaría el del proyecto.
    return create_turbine_layout_layer(
        coords=coords,
        layer_name=str(values.get("layer_name") or ""),
        module="energy",
        model_name="",
        coords_csv=coords_csv,
        crs_authid=_crs_definition(crs),
    )


__all__ = [
    "read_xy_csv",
    "read_xy_table",
    "create_turbine_layout_layer",
    "import_energy_layout",
    "import_turbine_layout_from_csv",
]

